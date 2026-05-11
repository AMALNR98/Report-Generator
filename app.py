from datetime import date
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

import pandas as pd
from flask import Flask, jsonify, render_template_string, request, send_file
from openpyxl import Workbook


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = BASE_DIR / "templates_excel"

REPORT_TYPES = {
    "monthly": {
        "label": "Monthly",
        "template": TEMPLATE_DIR / "monthly_template.xlsx",
        "filename_prefix": "Monthly_Alert_Report",
    },
    "weekly": {
        "label": "Weekly",
        "template": TEMPLATE_DIR / "weekly_template.xlsx",
        "filename_prefix": "Weekly_Alert_Report",
    },
    "daily": {
        "label": "Daily",
        "template": TEMPLATE_DIR / "daily_template.xlsx",
        "filename_prefix": "Daily_Alert_Report",
    },
}

ALLOWED_EXTENSIONS = {".xlsx", ".xls"}
OUTPUT_FIELDS = [
    "Date",
    "Alert Id",
    "Type of Alert",
    "Alert Level",
    "Severity",
    "Origin",
    "Entity",
    "Entity Type",
    "Status",
    "Message",
    "SLA (Y/N)",
]
EXCEL_COLUMNS = {
    "Date": 1,
    "Alert Id": 2,
    "Type of Alert": 5,
    "Alert Level": 8,
    "Severity": 11,
    "Origin": 12,
    "Entity": 14,
    "Entity Type": 16,
    "Status": 17,
    "Message": 18,
    "SLA (Y/N)": 19,
}
HEADER_ROW = 14
START_ROW = 15
SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
ET.register_namespace("", SPREADSHEET_NS)
ET.register_namespace("r", OFFICE_REL_NS)
REQUIRED_RAW_COLUMNS = [
    "Time",
    "Alert Id",
    "Type",
    "Severity",
    "Confidence",
    "Entity Type",
    "Status",
    "Message",
    "Assignee Email",
]


def today_string():
    return date.today().isoformat()


def clean_text(value):
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime().replace(tzinfo=None)
    if hasattr(value, "isoformat") and not isinstance(value, str):
        value = value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
    text = str(value)
    text = re.sub(r"[\ud800-\udfff]", "", text)
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ch == "\r" or ord(ch) >= 32)
    return text.strip()


def validate_report_type(report_type):
    key = clean_text(report_type).lower()
    if key not in REPORT_TYPES:
        raise ValueError("Select a valid report type: Monthly, Weekly, or Daily.")
    return key


def validate_upload(file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError("Please upload a raw alert Excel export.")
    suffix = Path(file_storage.filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Invalid file type. Upload a .xlsx or .xls raw alert export.")
    return suffix


def choose_column(columns, preferred, fallback=None):
    column_set = set(columns)
    if preferred in column_set:
        return preferred
    if fallback and fallback in column_set:
        return fallback
    return None


def choose_column_by_excel_letter(columns, letter):
    index = ord(letter.upper()) - ord("A")
    if 0 <= index < len(columns):
        return columns[index]
    return None


def choose_first_available(*values):
    for value in values:
        if value:
            return value
    return None


def parse_alert_file(file_storage):
    validate_upload(file_storage)
    data = BytesIO(file_storage.read())
    try:
        df = pd.read_excel(data)
    except Exception as exc:
        raise ValueError(f"Unable to read the Excel file. Check that it is a valid .xlsx or .xls export. Details: {exc}") from exc

    columns = [clean_text(col) for col in df.columns]
    df.columns = columns

    missing = [col for col in REQUIRED_RAW_COLUMNS if col not in df.columns]
    origin_col = choose_first_available(choose_column_by_excel_letter(df.columns, "H"), choose_column(df.columns, "Origin.1", "Origin"))
    entity_col = choose_first_available(choose_column_by_excel_letter(df.columns, "K"), choose_column(df.columns, "Entity.1", "Entity"))
    if origin_col is None:
        missing.append("Origin column H or Origin.1")
    if entity_col is None:
        missing.append("Entity column K or Entity.1")
    if missing:
        raise ValueError("Missing required raw alert columns: " + ", ".join(missing))

    rows = []
    for _, record in df.iterrows():
        assignee = clean_text(record.get("Assignee Email"))
        rows.append(
            {
                "Date": clean_text(record.get("Time")),
                "Alert Id": clean_text(record.get("Alert Id")),
                "Type of Alert": clean_text(record.get("Type")),
                "Alert Level": clean_text(record.get("Severity")),
                "Severity": clean_text(record.get("Confidence")),
                "Origin": clean_text(record.get(origin_col)),
                "Entity": clean_text(record.get(entity_col)),
                "Entity Type": clean_text(record.get("Entity Type")),
                "Status": clean_text(record.get("Status")),
                "Message": clean_text(record.get("Message")),
                "SLA (Y/N)": "Y" if "@" in assignee else "N",
            }
        )
    return rows


def compute_stats(rows):
    def norm(value):
        return clean_text(value).upper()

    return {
        "total": len(rows),
        "critical": sum(1 for row in rows if norm(row.get("Alert Level")) == "CRITICAL"),
        "major": sum(1 for row in rows if norm(row.get("Alert Level")) == "MAJOR"),
        "sla_y": sum(1 for row in rows if norm(row.get("SLA (Y/N)")) == "Y"),
    }


def sanitize_rows(rows):
    if not isinstance(rows, list):
        raise ValueError("Processed rows must be provided as a JSON array.")
    sanitized = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("Each processed row must be an object.")
        sanitized.append({field: clean_text(row.get(field, "")) for field in OUTPUT_FIELDS})
    return sanitized


def preview_workbook(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Alert Preview"
    worksheet.append(OUTPUT_FIELDS)
    for row in rows:
        worksheet.append([row.get(field, "") for field in OUTPUT_FIELDS])

    for column_cells in worksheet.columns:
        max_length = max(len(clean_text(cell.value)) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def ns_tag(name):
    return f"{{{SPREADSHEET_NS}}}{name}"


def rel_tag(name):
    return f"{{{REL_NS}}}{name}"


def office_rel_attr(name):
    return f"{{{OFFICE_REL_NS}}}{name}"


def col_num_from_ref(ref):
    match = re.match(r"([A-Z]+)", ref)
    if not match:
        return 0
    col_idx = 0
    for char in match.group(1):
        col_idx = col_idx * 26 + ord(char) - ord("A") + 1
    return col_idx


def col_letter(col_idx):
    letters = ""
    while col_idx:
        col_idx, remainder = divmod(col_idx - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def cell_ref(row_idx, col_idx):
    return f"{col_letter(col_idx)}{row_idx}"


def range_ref(min_row, min_col, max_row, max_col):
    start = cell_ref(min_row, min_col)
    end = cell_ref(max_row, max_col)
    return start if start == end else f"{start}:{end}"


def read_shared_strings(zip_file):
    if "xl/sharedStrings.xml" not in zip_file.namelist():
        return []
    root = ET.fromstring(zip_file.read("xl/sharedStrings.xml"))
    values = []
    for item in root.findall(ns_tag("si")):
        values.append("".join(text.text or "" for text in item.findall(f".//{ns_tag('t')}")))
    return values


def cell_text(cell, shared_strings):
    value = cell.find(ns_tag("v"))
    inline = cell.find(ns_tag("is"))
    if inline is not None:
        return "".join(text.text or "" for text in inline.findall(f".//{ns_tag('t')}"))
    if value is None:
        return ""
    text = value.text or ""
    if cell.attrib.get("t") == "s" and text.isdigit():
        index = int(text)
        if index < len(shared_strings):
            return shared_strings[index]
    return text


def workbook_sheet_paths(zip_file):
    workbook = ET.fromstring(zip_file.read("xl/workbook.xml"))
    relationships = ET.fromstring(zip_file.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships.findall(rel_tag("Relationship"))
    }
    sheets = []
    for sheet in workbook.find(ns_tag("sheets")).findall(ns_tag("sheet")):
        relationship_id = sheet.attrib[office_rel_attr("id")]
        target = rel_targets[relationship_id]
        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = "xl/" + target.lstrip("/")
        sheets.append({"name": sheet.attrib["name"], "path": path, "sheet_id": sheet.attrib.get("sheetId")})
    return sheets


def set_active_report_sheet(workbook_xml, report_index):
    workbook = ET.fromstring(workbook_xml)
    workbook_view = workbook.find(f"{ns_tag('bookViews')}/{ns_tag('workbookView')}")
    if workbook_view is not None:
        workbook_view.set("activeTab", str(report_index))
        workbook_view.set("firstSheet", "0")
    return ET.tostring(workbook, encoding="utf-8", xml_declaration=True)


def normalize_header(value):
    return re.sub(r"\s+", " ", clean_text(value)).strip().lower()


def hidden_columns(sheet_root):
    hidden = set()
    cols = sheet_root.find(ns_tag("cols"))
    if cols is None:
        return hidden
    for col in cols.findall(ns_tag("col")):
        if col.attrib.get("hidden") == "1":
            min_col = int(float(col.attrib.get("min", "0")))
            max_col = int(float(col.attrib.get("max", min_col)))
            hidden.update(range(min_col, max_col + 1))
    return hidden


def visible_column_in_range(hidden, min_col, max_col):
    for col_idx in range(min_col, max_col + 1):
        if col_idx not in hidden:
            return col_idx
    return min_col


def parse_range(ref):
    start, _, end = ref.partition(":")
    end = end or start
    start_match = re.match(r"([A-Z]+)(\d+)", start)
    end_match = re.match(r"([A-Z]+)(\d+)", end)
    if not start_match or not end_match:
        return None
    return {
        "min_col": col_num_from_ref(start),
        "min_row": int(start_match.group(2)),
        "max_col": col_num_from_ref(end),
        "max_row": int(end_match.group(2)),
    }


def merged_ranges(sheet_root):
    merge_cells = sheet_root.find(ns_tag("mergeCells"))
    if merge_cells is None:
        return []
    ranges = []
    for merge_cell in merge_cells.findall(ns_tag("mergeCell")):
        parsed = parse_range(merge_cell.attrib.get("ref", ""))
        if parsed:
            ranges.append(parsed)
    return ranges


def body_merge_patterns(ranges):
    patterns = []
    for merged_range in ranges:
        if merged_range["min_row"] <= START_ROW <= merged_range["max_row"]:
            patterns.append(
                {
                    "min_col": merged_range["min_col"],
                    "max_col": merged_range["max_col"],
                    "row_span": merged_range["max_row"] - merged_range["min_row"],
                }
            )
    return patterns


def ensure_body_merges(sheet_root, ranges, max_row):
    patterns = body_merge_patterns(ranges)
    if not patterns:
        return ranges

    merge_cells = sheet_root.find(ns_tag("mergeCells"))
    if merge_cells is None:
        merge_cells = ET.Element(ns_tag("mergeCells"))
        sheet_data = sheet_root.find(ns_tag("sheetData"))
        children = list(sheet_root)
        insert_at = children.index(sheet_data) + 1 if sheet_data in children else len(children)
        sheet_root.insert(insert_at, merge_cells)

    existing = {merge_cell.attrib.get("ref") for merge_cell in merge_cells.findall(ns_tag("mergeCell"))}
    for row_idx in range(START_ROW, max_row + 1):
        for pattern in patterns:
            ref = range_ref(row_idx, pattern["min_col"], row_idx + pattern["row_span"], pattern["max_col"])
            if ref not in existing:
                merge_cells.append(ET.Element(ns_tag("mergeCell"), {"ref": ref}))
                existing.add(ref)
                ranges.append(
                    {
                        "min_row": row_idx,
                        "max_row": row_idx + pattern["row_span"],
                        "min_col": pattern["min_col"],
                        "max_col": pattern["max_col"],
                    }
                )
    merge_cells.set("count", str(len(merge_cells.findall(ns_tag("mergeCell")))))
    return ranges


def writable_position(row_idx, col_idx, ranges):
    for merged_range in ranges:
        if (
            merged_range["min_row"] <= row_idx <= merged_range["max_row"]
            and merged_range["min_col"] <= col_idx <= merged_range["max_col"]
        ):
            return merged_range["min_row"], merged_range["min_col"]
    return row_idx, col_idx


def get_or_create_row(sheet_data, row_idx):
    rows = sheet_data.findall(ns_tag("row"))
    for row in rows:
        if int(row.attrib.get("r", "0")) == row_idx:
            return row

    row = ET.Element(ns_tag("row"), {"r": str(row_idx)})
    inserted = False
    for index, existing in enumerate(rows):
        if int(existing.attrib.get("r", "0")) > row_idx:
            sheet_data.insert(index, row)
            inserted = True
            break
    if not inserted:
        sheet_data.append(row)
    return row


def get_or_create_cell(row, row_idx, col_idx):
    ref = cell_ref(row_idx, col_idx)
    cells = row.findall(ns_tag("c"))
    for cell in cells:
        if cell.attrib.get("r") == ref:
            return cell

    cell = ET.Element(ns_tag("c"), {"r": ref})
    inserted = False
    for index, existing in enumerate(cells):
        if col_num_from_ref(existing.attrib.get("r", "")) > col_idx:
            row.insert(index, cell)
            inserted = True
            break
    if not inserted:
        row.append(cell)
    return cell


def row_style_template(sheet_data):
    template_row = get_or_create_row(sheet_data, START_ROW)
    row_attrs = {key: value for key, value in template_row.attrib.items() if key != "r"}
    cell_styles = {}
    for cell in template_row.findall(ns_tag("c")):
        col_idx = col_num_from_ref(cell.attrib.get("r", ""))
        if col_idx and "s" in cell.attrib:
            cell_styles[col_idx] = cell.attrib["s"]
    return row_attrs, cell_styles


def apply_row_style(row, row_idx, row_attrs):
    row.attrib.clear()
    row.attrib["r"] = str(row_idx)
    row.attrib.update(row_attrs)


def base_style_for_column(cell_styles, col_idx):
    if col_idx in cell_styles:
        return cell_styles[col_idx]
    lower = [idx for idx in cell_styles if idx <= col_idx]
    if lower:
        return cell_styles[max(lower)]
    return None


def set_cell_value(sheet_data, row_idx, col_idx, value, ranges, row_attrs=None, style_id=None):
    row_idx, col_idx = writable_position(row_idx, col_idx, ranges)
    row = get_or_create_row(sheet_data, row_idx)
    if row_attrs is not None:
        apply_row_style(row, row_idx, row_attrs)
    cell = get_or_create_cell(row, row_idx, col_idx)
    if style_id is not None:
        cell.set("s", str(style_id))
    for child in list(cell):
        if child.tag in {ns_tag("v"), ns_tag("is"), ns_tag("f")}:
            cell.remove(child)
    if value == "":
        cell.attrib.pop("t", None)
        return
    cell.set("t", "inlineStr")
    inline = ET.SubElement(cell, ns_tag("is"))
    text = ET.SubElement(inline, ns_tag("t"))
    text.set(f"{{{XML_NS}}}space", "preserve")
    text.text = clean_text(value)


def template_field_columns(sheet_root, shared_strings):
    field_columns = {}
    field_names = {normalize_header(field): field for field in OUTPUT_FIELDS}
    hidden = hidden_columns(sheet_root)
    ranges = merged_ranges(sheet_root)
    sheet_data = sheet_root.find(ns_tag("sheetData"))
    header_cells = {}

    for row in sheet_data.findall(ns_tag("row")):
        if int(row.attrib.get("r", "0")) == HEADER_ROW:
            header_cells = {col_num_from_ref(cell.attrib.get("r", "")): cell for cell in row.findall(ns_tag("c"))}
            break

    for merged_range in ranges:
        if merged_range["min_row"] <= HEADER_ROW <= merged_range["max_row"]:
            cell = header_cells.get(merged_range["min_col"])
            header_value = cell_text(cell, shared_strings) if cell is not None else ""
            field = field_names.get(normalize_header(header_value))
            if field:
                field_columns[field] = visible_column_in_range(hidden, merged_range["min_col"], merged_range["max_col"])

    for col_idx, cell in header_cells.items():
        field = field_names.get(normalize_header(cell_text(cell, shared_strings)))
        if field and field not in field_columns:
            field_columns[field] = col_idx

    resolved = EXCEL_COLUMNS.copy()
    resolved.update(field_columns)
    return resolved


def sheet_max_row(sheet_root):
    rows = sheet_root.find(ns_tag("sheetData")).findall(ns_tag("row"))
    if not rows:
        return 0
    return max(int(row.attrib.get("r", "0")) for row in rows)


def sheet_max_col(sheet_data):
    max_col = 0
    for row in sheet_data.findall(ns_tag("row")):
        for cell in row.findall(ns_tag("c")):
            max_col = max(max_col, col_num_from_ref(cell.attrib.get("r", "")))
    return max_col


def update_sheet_dimension(sheet_root, max_row, max_col):
    dimension = sheet_root.find(ns_tag("dimension"))
    if dimension is not None and max_row and max_col:
        dimension.set("ref", f"A1:{cell_ref(max_row, max_col)}")


def apply_body_row_format(sheet_data, row_idx, max_col, row_attrs, cell_styles, ranges):
    row = get_or_create_row(sheet_data, row_idx)
    apply_row_style(row, row_idx, row_attrs)
    for col_idx in range(1, max_col + 1):
        write_row, write_col = writable_position(row_idx, col_idx, ranges)
        if write_row != row_idx:
            continue
        cell = get_or_create_cell(row, row_idx, write_col)
        style_id = base_style_for_column(cell_styles, write_col)
        if style_id is not None:
            cell.set("s", str(style_id))


def style_tag(name):
    return f"{{{SPREADSHEET_NS}}}{name}"


def next_style_id(cell_xfs):
    return str(len(cell_xfs.findall(style_tag("xf"))))


def append_fill_style(styles_xml, base_style_id, rgb_color):
    styles_root = ET.fromstring(styles_xml)
    fills = styles_root.find(style_tag("fills"))
    cell_xfs = styles_root.find(style_tag("cellXfs"))
    if fills is None or cell_xfs is None:
        return styles_xml, base_style_id

    fill = ET.Element(style_tag("fill"))
    pattern_fill = ET.SubElement(fill, style_tag("patternFill"), {"patternType": "solid"})
    ET.SubElement(pattern_fill, style_tag("fgColor"), {"rgb": rgb_color})
    ET.SubElement(pattern_fill, style_tag("bgColor"), {"indexed": "64"})
    fills.append(fill)
    fill_id = str(len(fills.findall(style_tag("fill"))) - 1)
    fills.set("count", str(len(fills.findall(style_tag("fill")))))

    xfs = cell_xfs.findall(style_tag("xf"))
    base_index = int(base_style_id) if str(base_style_id).isdigit() else 0
    base_xf = deepcopy(xfs[base_index] if base_index < len(xfs) else xfs[0])
    base_xf.set("fillId", fill_id)
    base_xf.set("applyFill", "1")
    style_id = next_style_id(cell_xfs)
    cell_xfs.append(base_xf)
    cell_xfs.set("count", str(len(cell_xfs.findall(style_tag("xf")))))
    return ET.tostring(styles_root, encoding="utf-8", xml_declaration=True), style_id


def alert_level_style_map(styles_xml, base_style_id):
    styles_xml, critical_style = append_fill_style(styles_xml, base_style_id, "FFFF0000")
    styles_xml, major_style = append_fill_style(styles_xml, base_style_id, "FFFFC000")
    styles_xml, minor_style = append_fill_style(styles_xml, base_style_id, "FFFFFF00")
    return styles_xml, {
        "CRITICAL": critical_style,
        "MAJOR": major_style,
        "MINOR": minor_style,
    }


def filled_report_workbook(report_type, rows):
    if not rows:
        raise ValueError("No processed alert rows were provided. Process a raw alert file before generating the report.")

    config = REPORT_TYPES[report_type]
    template_path = config["template"]
    if not template_path.exists():
        raise ValueError(f"Missing local template file: {template_path.relative_to(BASE_DIR)}")

    output = BytesIO()
    try:
        with ZipFile(template_path, "r") as source:
            sheet_paths = workbook_sheet_paths(source)
            report_sheet = next((sheet for sheet in sheet_paths if sheet["name"] == "Report"), None)
            if report_sheet is None:
                raise ValueError('The selected template must contain a worksheet named "Report".')

            shared_strings = read_shared_strings(source)
            sheet_root = ET.fromstring(source.read(report_sheet["path"]))
            sheet_data = sheet_root.find(ns_tag("sheetData"))
            if sheet_data is None:
                raise ValueError('The selected template "Report" worksheet is missing sheet data.')

            ranges = merged_ranges(sheet_root)
            field_columns = template_field_columns(sheet_root, shared_strings)
            row_attrs, cell_styles = row_style_template(sheet_data)
            default_styles = {
                field: base_style_for_column(cell_styles, col_idx)
                for field, col_idx in field_columns.items()
            }
            alert_base_style = default_styles.get("Alert Level") or "0"
            updated_styles = None
            alert_styles = {}
            if "xl/styles.xml" in source.namelist():
                updated_styles, alert_styles = alert_level_style_map(source.read("xl/styles.xml"), alert_base_style)
            max_row = sheet_max_row(sheet_root)
            required_max_row = max(max_row, START_ROW + len(rows) - 1)
            max_col = max(sheet_max_col(sheet_data), max(field_columns.values()))
            ranges = ensure_body_merges(sheet_root, ranges, required_max_row)
            update_sheet_dimension(sheet_root, required_max_row, max_col)

            for row_idx in range(START_ROW, required_max_row + 1):
                apply_body_row_format(sheet_data, row_idx, max_col, row_attrs, cell_styles, ranges)
                for field, col_idx in field_columns.items():
                    set_cell_value(sheet_data, row_idx, col_idx, "", ranges, row_attrs, default_styles.get(field))

            for offset, row in enumerate(rows):
                excel_row = START_ROW + offset
                for field, col_idx in field_columns.items():
                    style_id = default_styles.get(field)
                    if field == "Alert Level":
                        style_id = alert_styles.get(clean_text(row.get(field, "")).upper(), style_id)
                    set_cell_value(sheet_data, excel_row, col_idx, row.get(field, ""), ranges, row_attrs, style_id)

            report_index = sheet_paths.index(report_sheet)
            updated_sheet = ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
            updated_workbook = set_active_report_sheet(source.read("xl/workbook.xml"), report_index)

            with ZipFile(output, "w", ZIP_DEFLATED) as destination:
                for item in source.infolist():
                    if item.filename == report_sheet["path"]:
                        destination.writestr(item, updated_sheet)
                    elif item.filename == "xl/workbook.xml":
                        destination.writestr(item, updated_workbook)
                    elif item.filename == "xl/styles.xml" and updated_styles is not None:
                        destination.writestr(item, updated_styles)
                    else:
                        destination.writestr(item, source.read(item.filename))
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Unable to fill the local template. Details: {exc}") from exc

    output.seek(0)
    return output


def excel_response(buffer, filename):
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/")
def index():
    template_status = {
        key: {
            "label": value["label"],
            "path": str(value["template"].relative_to(BASE_DIR)),
            "exists": value["template"].exists(),
        }
        for key, value in REPORT_TYPES.items()
    }
    return render_template_string(INDEX_HTML, template_status=template_status)


@app.post("/process")
def process():
    try:
        report_type = validate_report_type(request.form.get("report_type", ""))
        rows = parse_alert_file(request.files.get("raw_file"))
        return jsonify(
            {
                "ok": True,
                "report_type": report_type,
                "report_label": REPORT_TYPES[report_type]["label"],
                "rows": rows,
                "count": len(rows),
                "stats": compute_stats(rows),
                "template": {
                    "path": str(REPORT_TYPES[report_type]["template"].relative_to(BASE_DIR)),
                    "exists": REPORT_TYPES[report_type]["template"].exists(),
                },
            }
        )
    except ValueError as exc:
        return jsonify({"ok": False, "error": clean_text(exc)}), 400


@app.post("/download_preview")
def download_preview():
    try:
        payload = request.get_json(silent=True) or {}
        rows = sanitize_rows(payload.get("rows", []))
        return excel_response(preview_workbook(rows), f"Alert_Preview_{today_string()}.xlsx")
    except ValueError as exc:
        return jsonify({"ok": False, "error": clean_text(exc)}), 400


@app.post("/download")
def download():
    try:
        payload = request.get_json(silent=True) or {}
        report_type = validate_report_type(payload.get("report_type", ""))
        rows = sanitize_rows(payload.get("rows", []))
        filename = f"{REPORT_TYPES[report_type]['filename_prefix']}_{today_string()}.xlsx"
        return excel_response(filled_report_workbook(report_type, rows), filename)
    except ValueError as exc:
        return jsonify({"ok": False, "error": clean_text(exc)}), 400


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"ok": False, "error": "The uploaded file is too large. Maximum size is 50 MB."}), 413


INDEX_HTML = """
<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Monthly, Weekly, Daily Alert Report Generator</title>
  <style>
    :root {
      color-scheme: dark;
      --bg: #0f1720;
      --panel: #151f2b;
      --panel-2: #1b2634;
      --line: #2c3b4f;
      --text: #e7edf5;
      --muted: #9aacbd;
      --accent: #28a7c7;
      --accent-2: #8bd450;
      --danger: #ff5d66;
      --warn: #ffb24a;
      --minor: #57d68d;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
      min-height: 100vh;
    }
    .app {
      display: grid;
      grid-template-columns: minmax(300px, 360px) minmax(0, 1fr);
      min-height: 100vh;
    }
    aside {
      border-right: 1px solid var(--line);
      background: #121c27;
      padding: 24px;
      overflow: auto;
    }
    main {
      padding: 24px;
      overflow: auto;
    }
    h1 {
      font-size: 22px;
      line-height: 1.2;
      margin: 0 0 8px;
      letter-spacing: 0;
    }
    h2 {
      font-size: 15px;
      margin: 22px 0 10px;
      color: #cdd9e6;
      letter-spacing: 0;
    }
    p, li, label, small { color: var(--muted); }
    .subtitle { margin: 0 0 18px; }
    .group {
      border-top: 1px solid var(--line);
      padding-top: 18px;
      margin-top: 18px;
    }
    .segmented {
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 8px;
    }
    .segmented input { position: absolute; opacity: 0; pointer-events: none; }
    .segmented span {
      display: block;
      text-align: center;
      border: 1px solid var(--line);
      background: var(--panel);
      padding: 10px 8px;
      border-radius: 6px;
      color: var(--text);
      cursor: pointer;
      font-weight: 700;
    }
    .segmented input:checked + span {
      background: #103244;
      border-color: var(--accent);
      color: #d9f8ff;
    }
    .dropzone {
      border: 1px dashed #52708d;
      background: var(--panel);
      border-radius: 8px;
      padding: 18px;
      text-align: center;
      transition: border-color .15s ease, background .15s ease;
    }
    .dropzone.dragover {
      border-color: var(--accent);
      background: #132b38;
    }
    input[type="file"] { display: none; }
    .file-button, button {
      appearance: none;
      border: 1px solid transparent;
      border-radius: 6px;
      padding: 10px 12px;
      min-height: 42px;
      font-weight: 800;
      color: #061019;
      background: var(--accent);
      cursor: pointer;
      width: 100%;
      margin-top: 10px;
    }
    button.secondary {
      background: #263649;
      color: var(--text);
      border-color: #3b4c62;
    }
    button.success {
      background: var(--accent-2);
    }
    button:disabled {
      cursor: not-allowed;
      opacity: .45;
    }
    .mapping {
      display: grid;
      gap: 6px;
      margin: 0;
      padding: 0;
      list-style: none;
      font-size: 13px;
    }
    .mapping li {
      display: flex;
      justify-content: space-between;
      gap: 10px;
      border-bottom: 1px solid rgba(255,255,255,.06);
      padding-bottom: 6px;
    }
    .mapping strong { color: var(--text); text-align: right; }
    .status {
      margin-top: 12px;
      padding: 10px 12px;
      border-radius: 6px;
      background: var(--panel);
      color: var(--muted);
      border: 1px solid var(--line);
      min-height: 42px;
      line-height: 1.35;
    }
    .status.ok { color: #c8f7ce; border-color: #326b44; }
    .status.error { color: #ffd4d7; border-color: #8a3037; }
    .template-pill {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 10px;
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 10px 12px;
      font-size: 13px;
    }
    .dot {
      width: 10px;
      height: 10px;
      border-radius: 50%;
      flex: 0 0 auto;
      background: var(--danger);
    }
    .dot.ok { background: var(--accent-2); }
    .stats {
      display: grid;
      grid-template-columns: repeat(4, minmax(120px, 1fr));
      gap: 12px;
      margin-bottom: 18px;
    }
    .stat {
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 14px;
    }
    .stat b { display: block; font-size: 28px; color: var(--text); }
    .stat span { color: var(--muted); font-size: 13px; }
    .table-wrap {
      border: 1px solid var(--line);
      border-radius: 8px;
      overflow: auto;
      max-height: calc(100vh - 190px);
      background: var(--panel);
    }
    table {
      width: 100%;
      border-collapse: collapse;
      min-width: 1200px;
    }
    th, td {
      padding: 10px 12px;
      border-bottom: 1px solid rgba(255,255,255,.07);
      text-align: left;
      vertical-align: top;
      font-size: 13px;
    }
    th {
      position: sticky;
      top: 0;
      z-index: 2;
      background: #203044;
      color: #f3f7fb;
    }
    tr:hover td { background: rgba(255,255,255,.03); }
    .badge {
      display: inline-block;
      border-radius: 999px;
      padding: 3px 8px;
      font-weight: 800;
      font-size: 12px;
    }
    .critical { background: rgba(255,93,102,.16); color: #ff9aa0; }
    .major { background: rgba(255,178,74,.16); color: #ffc979; }
    .minor { background: rgba(87,214,141,.16); color: #8ff0b5; }
    .sla-y { color: #8ff0b5; font-weight: 900; }
    .sla-n { color: #ff9aa0; font-weight: 900; }
    .empty {
      color: var(--muted);
      padding: 34px;
      text-align: center;
    }
    @media (max-width: 980px) {
      .app { grid-template-columns: 1fr; }
      aside { border-right: 0; border-bottom: 1px solid var(--line); }
      .stats { grid-template-columns: repeat(2, 1fr); }
    }
  </style>
</head>
<body>
  <div class="app">
    <aside>
      <h1>Monthly, Weekly, Daily Alert Report Generator</h1>
      <p class="subtitle">Local Excel conversion using project templates only.</p>

      <form id="processForm">
        <div class="group">
          <h2>Report Type</h2>
          <div class="segmented" role="radiogroup" aria-label="Report type">
            <label><input type="radio" name="report_type" value="monthly" checked><span>Monthly</span></label>
            <label><input type="radio" name="report_type" value="weekly"><span>Weekly</span></label>
            <label><input type="radio" name="report_type" value="daily"><span>Daily</span></label>
          </div>
        </div>

        <div class="group">
          <h2>Raw Alert Upload</h2>
          <div id="dropzone" class="dropzone">
            <p id="fileName">Drop a .xlsx or .xls raw alert export here</p>
            <label class="file-button" for="rawFile">Choose Excel File</label>
            <input id="rawFile" name="raw_file" type="file" accept=".xlsx,.xls">
          </div>
        </div>

        <div class="group">
          <h2>Local Template</h2>
          <div class="template-pill">
            <span id="templatePath"></span>
            <span id="templateDot" class="dot" aria-hidden="true"></span>
          </div>
        </div>

        <div class="group">
          <h2>Save Location</h2>
          <div class="template-pill">
            <span id="saveCapability"></span>
          </div>
        </div>

        <div class="group">
          <h2>Field Mapping</h2>
          <ul class="mapping">
            <li><span>Time</span><strong>Date</strong></li>
            <li><span>Alert Id</span><strong>Alert Id</strong></li>
            <li><span>Type</span><strong>Type of Alert</strong></li>
            <li><span>Severity</span><strong>Alert Level</strong></li>
            <li><span>Confidence</span><strong>Severity</strong></li>
            <li><span>Origin (2nd col)</span><strong>Origin</strong></li>
            <li><span>Entity (col K)</span><strong>Entity</strong></li>
            <li><span>Entity Type</span><strong>Entity Type</strong></li>
            <li><span>Status</span><strong>Status</strong></li>
            <li><span>Message</span><strong>Message</strong></li>
            <li><span>Assignee Email</span><strong>SLA (Y/N)</strong></li>
          </ul>
        </div>

        <div class="group">
          <button id="processBtn" type="submit">Process & Preview</button>
          <button id="saveBtn" class="success" type="button" disabled>Save Report to Folder</button>
          <button id="previewBtn" class="secondary" type="button" disabled>Download Excel Preview</button>
          <button id="reportBtn" class="secondary" type="button" disabled>Download Generated Report</button>
        </div>
        <div id="status" class="status">Ready.</div>
      </form>
    </aside>

    <main>
      <div class="stats">
        <div class="stat"><b id="totalStat">0</b><span>Total records</span></div>
        <div class="stat"><b id="criticalStat">0</b><span>Critical count</span></div>
        <div class="stat"><b id="majorStat">0</b><span>Major count</span></div>
        <div class="stat"><b id="slaStat">0</b><span>SLA Y count</span></div>
      </div>
      <div class="table-wrap">
        <table id="previewTable">
          <thead></thead>
          <tbody><tr><td class="empty">Process a raw alert export to preview records.</td></tr></tbody>
        </table>
      </div>
    </main>
  </div>

  <script>
    const templateStatus = {{ template_status | tojson }};
    const fields = ["Date", "Alert Id", "Type of Alert", "Alert Level", "Severity", "Origin", "Entity", "Entity Type", "Status", "Message", "SLA (Y/N)"];
    let processedRows = [];
    let currentReportType = "monthly";

    const form = document.getElementById("processForm");
    const rawFile = document.getElementById("rawFile");
    const dropzone = document.getElementById("dropzone");
    const fileName = document.getElementById("fileName");
    const statusBox = document.getElementById("status");
    const templatePath = document.getElementById("templatePath");
    const templateDot = document.getElementById("templateDot");
    const saveCapability = document.getElementById("saveCapability");
    const processBtn = document.getElementById("processBtn");
    const saveBtn = document.getElementById("saveBtn");
    const previewBtn = document.getElementById("previewBtn");
    const reportBtn = document.getElementById("reportBtn");

    function selectedReportType() {
      return new FormData(form).get("report_type");
    }

    function setStatus(message, type = "") {
      statusBox.textContent = message;
      statusBox.className = "status" + (type ? " " + type : "");
    }

    function updateTemplateStatus() {
      currentReportType = selectedReportType();
      const info = templateStatus[currentReportType];
      templatePath.textContent = info.path + (info.exists ? " found" : " missing");
      templateDot.className = "dot" + (info.exists ? " ok" : "");
    }

    function updateButtons(enabled) {
      saveBtn.disabled = !enabled;
      previewBtn.disabled = !enabled;
      reportBtn.disabled = !enabled;
    }

    function validExcelFile(file) {
      return file && /\\.(xlsx|xls)$/i.test(file.name);
    }

    function escapeHtml(value) {
      return String(value ?? "").replace(/[&<>"']/g, char => ({
        "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#039;"
      }[char]));
    }

    function badgeForLevel(value) {
      const normalized = String(value || "").toUpperCase();
      if (normalized === "CRITICAL") return `<span class="badge critical">${escapeHtml(value)}</span>`;
      if (normalized === "MAJOR") return `<span class="badge major">${escapeHtml(value)}</span>`;
      if (normalized === "MINOR") return `<span class="badge minor">${escapeHtml(value)}</span>`;
      return escapeHtml(value);
    }

    function slaCell(value) {
      const normalized = String(value || "").toUpperCase();
      const cls = normalized === "Y" ? "sla-y" : "sla-n";
      return `<span class="${cls}">${escapeHtml(value)}</span>`;
    }

    function renderPreview(rows) {
      const table = document.getElementById("previewTable");
      table.querySelector("thead").innerHTML = `<tr>${fields.map(field => `<th>${escapeHtml(field)}</th>`).join("")}</tr>`;
      if (!rows.length) {
        table.querySelector("tbody").innerHTML = `<tr><td class="empty" colspan="${fields.length}">No records found.</td></tr>`;
        return;
      }
      table.querySelector("tbody").innerHTML = rows.map(row => {
        return `<tr>${fields.map(field => {
          if (field === "Alert Level") return `<td>${badgeForLevel(row[field])}</td>`;
          if (field === "SLA (Y/N)") return `<td>${slaCell(row[field])}</td>`;
          return `<td>${escapeHtml(row[field])}</td>`;
        }).join("")}</tr>`;
      }).join("");
    }

    function updateStats(stats) {
      document.getElementById("totalStat").textContent = stats.total || 0;
      document.getElementById("criticalStat").textContent = stats.critical || 0;
      document.getElementById("majorStat").textContent = stats.major || 0;
      document.getElementById("slaStat").textContent = stats.sla_y || 0;
    }

    async function downloadBlob(url, body, preferredName, saveToFolder = false) {
      const response = await fetch(url, {
        method: "POST",
        headers: { "Content-Type": "application/json" },
        body: JSON.stringify(body)
      });
      if (!response.ok) {
        const errorPayload = await response.json().catch(() => ({}));
        throw new Error(errorPayload.error || "The download request failed.");
      }
      const blob = await response.blob();
      const disposition = response.headers.get("Content-Disposition") || "";
      const match = disposition.match(/filename="?([^"]+)"?/i);
      const filename = match ? match[1] : preferredName;

      if (saveToFolder && window.showSaveFilePicker) {
        const handle = await window.showSaveFilePicker({
          suggestedName: filename,
          types: [{ description: "Excel workbook", accept: { "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": [".xlsx"] } }]
        });
        const writable = await handle.createWritable();
        await writable.write(blob);
        await writable.close();
        return filename;
      }

      const link = document.createElement("a");
      link.href = URL.createObjectURL(blob);
      link.download = filename;
      document.body.appendChild(link);
      link.click();
      link.remove();
      URL.revokeObjectURL(link.href);
      return filename;
    }

    form.addEventListener("change", event => {
      if (event.target.name === "report_type") {
        updateTemplateStatus();
        currentReportType = selectedReportType();
      }
    });

    rawFile.addEventListener("change", () => {
      const file = rawFile.files[0];
      fileName.textContent = file ? file.name : "Drop a .xlsx or .xls raw alert export here";
      if (file && !validExcelFile(file)) setStatus("Invalid file type. Upload .xlsx or .xls.", "error");
    });

    ["dragenter", "dragover"].forEach(eventName => {
      dropzone.addEventListener(eventName, event => {
        event.preventDefault();
        dropzone.classList.add("dragover");
      });
    });
    ["dragleave", "drop"].forEach(eventName => {
      dropzone.addEventListener(eventName, event => {
        event.preventDefault();
        dropzone.classList.remove("dragover");
      });
    });
    dropzone.addEventListener("drop", event => {
      const file = event.dataTransfer.files[0];
      if (!validExcelFile(file)) {
        setStatus("Invalid file type. Upload .xlsx or .xls.", "error");
        return;
      }
      rawFile.files = event.dataTransfer.files;
      fileName.textContent = file.name;
    });

    form.addEventListener("submit", async event => {
      event.preventDefault();
      const file = rawFile.files[0];
      if (!validExcelFile(file)) {
        setStatus("Choose a valid .xlsx or .xls raw alert export.", "error");
        return;
      }
      processBtn.disabled = true;
      updateButtons(false);
      setStatus("Processing raw alert export...");
      const formData = new FormData();
      formData.append("report_type", selectedReportType());
      formData.append("raw_file", file);
      try {
        const response = await fetch("/process", { method: "POST", body: formData });
        const payload = await response.json();
        if (!response.ok || !payload.ok) throw new Error(payload.error || "Processing failed.");
        processedRows = payload.rows;
        currentReportType = payload.report_type;
        updateStats(payload.stats);
        renderPreview(processedRows);
        updateButtons(true);
        setStatus(`${payload.report_label} preview ready with ${payload.count} records.`, "ok");
      } catch (error) {
        processedRows = [];
        updateStats({});
        renderPreview([]);
        setStatus(error.message, "error");
      } finally {
        processBtn.disabled = false;
      }
    });

    previewBtn.addEventListener("click", async () => {
      try {
        setStatus("Preparing preview workbook...");
        const filename = await downloadBlob("/download_preview", { rows: processedRows }, "Alert_Preview.xlsx");
        setStatus(`Preview ready: ${filename}`, "ok");
      } catch (error) {
        setStatus(error.message, "error");
      }
    });

    reportBtn.addEventListener("click", async () => {
      try {
        setStatus("Generating report from the selected local template...");
        const filename = await downloadBlob("/download", { report_type: currentReportType, rows: processedRows }, "Alert_Report.xlsx");
        setStatus(`Report ready: ${filename}`, "ok");
      } catch (error) {
        setStatus(error.message, "error");
      }
    });

    saveBtn.addEventListener("click", async () => {
      try {
        setStatus("Preparing report save...");
        const filename = await downloadBlob(
          "/download",
          { report_type: currentReportType, rows: processedRows },
          "Alert_Report.xlsx",
          Boolean(window.showSaveFilePicker)
        );
        setStatus(window.showSaveFilePicker ? `Saved report: ${filename}` : `Downloaded report: ${filename}`, "ok");
      } catch (error) {
        if (error.name === "AbortError") {
          setStatus("Save cancelled.");
        } else {
          setStatus(error.message, "error");
        }
      }
    });

    saveCapability.textContent = window.showSaveFilePicker
      ? "Folder save supported by this browser"
      : "Browser download fallback will be used";
    updateTemplateStatus();
  </script>
</body>
</html>
"""


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
